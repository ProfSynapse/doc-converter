"""Offline regression checks; Google services and the browser are mocked."""
import base64
from io import BytesIO
from pathlib import Path
from unittest.mock import MagicMock, patch

import pytest
from PIL import Image
from pypdf import PdfReader

from app.converters.google_docs_converter import GoogleDocsConverter
from app.converters.html_converter import HtmlConverter
from app.converters.markdown_converter import MarkdownConverter
from app.utils.render_resources import inline_raster_fetcher

FIXTURE = (Path(__file__).parent / 'fixtures/document.md').read_text()


def raster_url():
    data = BytesIO()
    Image.new('RGB', (2, 2), 'red').save(data, format='PNG')
    return 'data:image/png;base64,' + base64.b64encode(data.getvalue()).decode()


@pytest.mark.parametrize('url', [
    'https://example.com/image.png', 'http://127.0.0.1/image',
    'http://169.254.169.254/latest/meta-data', 'file:///etc/passwd',
    'data:image/svg+xml;base64,PHN2Zy8+',
    'data:image/x-eps;base64,JSFQUy1BZG9iZQ==',
    'data:image/png;base64,JSFQUy1BZG9iZQ==',
])
def test_pdf_asset_policy_rejects_non_raster_and_external_resources(url):
    with patch('socket.create_connection', side_effect=AssertionError('network used')):
        with pytest.raises(ValueError):
            inline_raster_fetcher(url)


def test_inline_raster_is_supported():
    result = inline_raster_fetcher(raster_url())
    assert result['mime_type'] == 'image/png'
    assert result['string'].startswith(b'\x89PNG')


def test_markdown_pdf_semantics_and_images(tmp_path):
    output = tmp_path / 'document.pdf'
    converter = MarkdownConverter()
    converter.convert_to_pdf(FIXTURE + '\n![Inline](' + raster_url() + ')', str(output))
    reader = PdfReader(output)
    text = '\n'.join(page.extract_text() for page in reader.pages)
    assert 'Security conversion fixture' in text
    assert 'Apples' in text and 'fixture' in text
    assert sum(len(page.images) for page in reader.pages) == 1


def test_markdown_pdf_never_falls_back_to_network(tmp_path):
    with patch('socket.create_connection', side_effect=AssertionError('network used')) as connect:
        MarkdownConverter().convert_to_pdf(
            '# Offline\n![No network](https://example.com/a.png)', str(tmp_path / 'offline.pdf'))
    connect.assert_not_called()


def test_markdown_docx_still_contains_heading_and_table(tmp_path):
    from docx import Document
    output = tmp_path / 'document.docx'
    MarkdownConverter().convert_to_docx(FIXTURE, str(output))
    doc = Document(output)
    assert any('Security conversion fixture' in p.text for p in doc.paragraphs)
    assert any('Apples' in cell.text for table in doc.tables for row in table.rows for cell in row.cells)


def test_html_pdf_disables_scripts_and_blocks_requests(tmp_path):
    output = tmp_path / 'html.pdf'
    manager = MagicMock()
    browser = manager.__enter__.return_value.chromium.launch.return_value
    context = browser.new_context.return_value
    context.new_page.return_value.pdf.side_effect = lambda **kw: Path(kw['path']).write_bytes(b'%PDF-mock')
    with patch('app.converters.html_converter.sync_playwright', return_value=manager):
        HtmlConverter().convert_to_pdf('<h1>Hello</h1><script>fetch("http://127.0.0.1")</script>', str(output))
    browser.new_context.assert_called_once_with(java_script_enabled=False, service_workers='block')
    route = MagicMock()
    pattern, handler = context.route.call_args.args
    assert pattern == '**/*'
    handler(route)
    route.abort.assert_called_once()


@pytest.mark.parametrize('html', [False, True])
def test_google_export_requests_and_private_default(html):
    docs, drive = MagicMock(), MagicMock()
    docs.documents.return_value.create.return_value.execute.return_value = {'documentId': 'fixture-id'}
    drive.files.return_value.get.return_value.execute.return_value = {'webViewLink': 'https://docs.google.com/document/d/fixture-id/edit'}
    converter = GoogleDocsConverter(docs, drive)
    if html:
        result = converter.convert_html('<h1>Fixture heading</h1><p>Hello <strong>world</strong></p>', 'Fixture')
    else:
        result = converter.convert('# Fixture heading\nHello **world**', 'Fixture')
    assert result['documentId'] == 'fixture-id'
    batch = docs.documents.return_value.batchUpdate.call_args.kwargs['body']['requests']
    assert any('Fixture heading' in r.get('insertText', {}).get('text', '') for r in batch)
    assert any('updateParagraphStyle' in r for r in batch)
    drive.permissions.assert_not_called()


def test_flask_factory_health_and_cors_use_isolated_configuration(monkeypatch, tmp_path):
    # Explicit dummy config prevents reading deployment credentials or opening a live DB.
    monkeypatch.setenv('SECRET_KEY', 'fixture-only-not-a-deployment-secret')
    monkeypatch.setenv('DATABASE_URL', 'sqlite:///:memory:')
    monkeypatch.setenv('CONVERTED_FOLDER', str(tmp_path / 'converted'))
    monkeypatch.setenv('FRONTEND_URL', 'https://converter.example.test')
    monkeypatch.delenv('GOOGLE_OAUTH_CLIENT_ID', raising=False)
    monkeypatch.delenv('GOOGLE_OAUTH_CLIENT_SECRET', raising=False)
    from app import create_app
    with patch('threading.Thread.start'), patch('app.utils.template_generator.ensure_template_exists', return_value=None):
        app = create_app('testing')
    client = app.test_client()
    response = client.get('/health', headers={'Origin': 'https://converter.example.test'})
    assert response.status_code == 200
    assert response.json['dependencies']['weasyprint'] == '70.0'
    assert response.headers['Access-Control-Allow-Origin'] == 'https://converter.example.test'
    denied = client.get('/health', headers={'Origin': 'https://other.example.test'})
    assert 'Access-Control-Allow-Origin' not in denied.headers
